from openpyxl import *
from openpyxl.styles import *

sv_xl = Workbook()
sv = sv_xl.active

"""установка линий ячеек"""
border = Side(color="000000", border_style="thin")
t = Border(top=border)
tb = Border(top=border, bottom=border)
lt = Border(left=border, top=border)
rt = Border(right=border, top=border)
ltb = Border(left=border, top=border, bottom=border)
ltr = Border(left=border, top=border, right=border)
ltrb = Border(left=border, top=border, right=border, bottom=border)

"""отрисовка таблицы линиями (A1:N25)"""
borders_cells = (  lt,   t,   t,  lt,   t,   t,   t,   t,   t,   t,   t,   t,   t,  rt,
                   lt,   t,   t,  lt,   t,   t,   t,   t,   t,  lt,   t,   t,   t,  rt,
                   lt,   t,   t,  lt,   t,   t,   t,   t,   t,  lt,   t,   t,   t,  rt,
                   lt,   t,   t,  lt,   t,   t,   t,   t,   t,  lt,   t,   t,   t,  rt,
                   lt,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,  rt,
                   lt,  lt,   t,  lt,   t,  lt,   t,  lt,   t,  lt,   t,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,  rt,
                   lt,  lt,   t,  lt,   t,  lt,   t,  lt,   t,  lt,   t,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,   t,  rt,
                   lt,  lt,   t,  lt,   t,  lt,   t,  lt,   t,  lt,   t,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                   lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt,  lt, ltr,
                  ltb, ltb, ltb, ltb, ltb, ltb, ltb, ltb, ltb, ltb, ltb, ltb, ltb, ltrb)

"""цикл отрисовки и установка размеров строк (20 ед.) и столбцов (10 ед.)"""
i = 0
for r in range(1, 26):
    for c in "abcdefghijklmn":
        if i <= len(borders_cells) - 1:
            sv[f"{c}{r}"].border = borders_cells[i]
            i += 1
        sv.column_dimensions[c].width = 10
        sv.row_dimensions[r].height = 20


"""установка в ячейку текста и его выравнивание"""
text_cells = [
    # cell, value, alignment
    # "", "", None
    # пример
    # "A1", "text", Alignment(vertical="center", horizontal="center")
    ["A1", "ДАТА", None],
    ["D1", "БРИГАДА", None],
    ["A2", "Распилено пиловочника", None],
    ["D2", "штук", Alignment(horizontal="center")]
]

"""цикл установки в ячейку текста и его выравнивание"""
i = 0
fontName = "Calibri"
size = 11
while i <= len(text_cells) - 1:
    sv[f"{text_cells[i][0]}"] = f"{text_cells[i][1]}"
    sv[f"{text_cells[i][0]}"].font = Font(name=fontName, size=size)
    if text_cells[i][2]:
        sv[f"{text_cells[i][0]}"].alignment = text_cells[i][2]
    i += 1

sv.merge_cells("E1:N1")
sv.merge_cells("D2:I2")
#sv["a1"] = "ДАТА"
#sv["a1"].font = Font(name="Calibri", size=14)
## horizontal= 'centerContinuous', 'distributed', 'general', 'center', 'left', 'justify', 'right', 'fill'
## vertical= 'top', 'justify', 'distributed', 'bottom', 'center'
#sv["a1"].alignment = Alignment(vertical=None, horizontal=None)



"""border_side = Side(border_style="hair")
border = Border(top=border_side,
                right=border_side,
                bottom=border_side,
                left=border_side)
"""
"""sv["a1"] = "ДАТА"
sv["a1"].border = Border(left=border_side, top=border_side)
sv.merge_cells("c1:d1")
sv["c1"].border = Border(top=border_side)
sv["d1"].border = Border(right=border_side)
sv["e1"] = "БРИГАДА"
sv["e1"].border = Border(top=border_side)
sv.merge_cells("g1:o1")
sv["f1"].border = Border(top=border_side)
sv["g1"].border = Border(top=border_side)
sv["o1"].border = Border(right=border_side)

sv["b3"] = "Распилено пиловочника"
sv["b3"].border = Border(left=border_side, top=border_side)
sv["c3"].border = Border(top=border_side)
sv["d3"].border = Border(top=border_side)
sv.merge_cells("e3:j3")
sv["e3"] = "штук"
sv["e3"].alignment = Alignment(horizontal="center")
sv["e3"].border = Border(left=border_side, top=border_side)
sv.merge_cells("k3:o3")
sv["k3"] = "всего м\u00B3"
sv["k3"].alignment = Alignment(horizontal="center")
sv["k3"].border = Border(left=border_side, top=border_side)
sv["o3"].border = Border(right=border_side)

sv["b4"] = "диаметр"
sv["b4"].border = Border(left=border_side, top=border_side)
sv["c4"].border = Border(top=border_side)
sv["d4"].border = Border(top=border_side)
sv.merge_cells("e4:j4")
sv["e4"].border = Border(left=border_side, top=border_side)
sv.merge_cells("k4:o4")
sv["k4"].border = Border(left=border_side, top=border_side)
sv["o4"].border = Border(right=border_side)

sv["b5"] = "диаметр"
sv["b5"].border = Border(left=border_side, top=border_side)
sv["c5"].border = Border(top=border_side)
sv["d5"].border = Border(top=border_side)
sv.merge_cells("e5:j5")
sv["e5"].border = Border(left=border_side, top=border_side)
sv.merge_cells("k5:o5")
sv["k5"].border = Border(left=border_side, top=border_side)
sv["o5"].border = Border(right=border_side)

for i in range(6, 21, 7):
    sv[f"b{i}"].border = Border(left=border_side, top=border_side)
    sv[f"c{i}"] = "НАПИЛЕНО ДОСКИ"
    sv[f"c{i}"].border = Border(top=border_side)
    sv[f"d{i}"].border = Border(top=border_side)
    sv.merge_cells(f"e{i}:f{i}")
    sv[f"e{i}"].border = Border(top=border_side)
    for w in ["g", "h", "i", "j", "k", "l", "m", "n", "o"]:
        sv[f"{w}{i}"].border = Border(top=border_side)
    sv[f"o{i}"].border = Border(top=border_side, right=border_side)

for i in range(7, 22, 7):
    sv[f"b{i}"].border = Border(left=border_side, top=border_side)
    sv.merge_cells(f"c{i}:d{i}")
    sv[f"c{i}"] = "0-4 с."
    sv[f"c{i}"].alignment = Alignment(horizontal="center")
    sv[f"c{i}"].border = Border(left=border_side, top=border_side)
    sv.merge_cells(f"e{i}:f{i}")
    sv[f"e{i}"] = "5 с."
    sv[f"e{i}"].alignment = Alignment(horizontal="center")
    sv[f"e{i}"].border = Border(left=border_side, top=border_side)
    sv.merge_cells(f"g{i}:h{i}")
    sv[f"g{i}"] = "6 с."
    sv[f"g{i}"].alignment = Alignment(horizontal="center")
    sv[f"g{i}"].border = Border(left=border_side, top=border_side)
    sv.merge_cells(f"i{i}:j{i}")
    sv[f"i{i}"] = "7 с."
    sv[f"i{i}"].alignment = Alignment(horizontal="center")
    sv[f"i{i}"].border = Border(left=border_side, top=border_side)
    sv.merge_cells(f"k{i}:l{i}")
    sv[f"k{i}"] = "некондиция"
    sv[f"k{i}"].alignment = Alignment(horizontal="center")
    sv[f"k{i}"].border = Border(left=border_side, top=border_side)
    sv[f"m{i}"]= "всего"
    sv[f"m{i}"].border = Border(left=border_side, top=border_side)
    sv[f"n{i}"] = "общая"
    sv[f"n{i}"].border = Border(left=border_side, top=border_side)
    sv[f"o{i}"] = "%"
    sv[f"o{i}"].border = Border(left=border_side, top=border_side, right=border_side)

for i in range(8, 23, 7):
    sv[f"b{i}"].border =Border(left=border_side, top=border_side)
    for w in ["cd", "ef", "gh", "ij", "kl", "mn"]:
        sv[f"{w[0]}{i}"] = "шт."
        sv[f"{w[0]}{i}"].border = Border(left=border_side, top=border_side)
        sv[f"{w[1]}{i}"] = "м\u00B3"
        sv[f"{w[1]}{i}"].border = Border(left=border_side, top=border_side)
    sv[f"o{i}"] = "выхода"
    sv[f"o{i}"].border = Border(left=border_side, top=border_side, right=border_side)

for i in range(9, 24, 7):
    sv[f"b{i}"] = "4 м"
    sv[f"b{i}"].alignment = Alignment(horizontal="center")
    for w in "bcdefghijklmno":
        sv[f"{w}{i}"].border = Border(left=border_side, top=border_side)
    sv[f"o{i}"].border = Border(left=border_side, top=border_side, right=border_side)

for i in range(10, 25, 7):
    sv[f"b{i}"] = "3 м"
    sv[f"b{i}"].alignment = Alignment(horizontal="center")
    for w in "bcdefghijklmno":
        sv[f"{w}{i}"].border = Border(left=border_side, top=border_side)
    sv[f"o{i}"].border = Border(left=border_side, top=border_side, right=border_side)

for i in range(11, 26, 7):
    sv[f"b{i}"] = "2 м"
    sv[f"b{i}"].alignment = Alignment(horizontal="center")
    for w in "bcdefghijklmno":
        sv[f"{w}{i}"].border = Border(left=border_side, top=border_side)
    sv[f"o{i}"].border = Border(left=border_side, top=border_side, right=border_side)

for i in range(12, 27, 7):
    sv[f"b{i}"] = "Итого:"
    for w in "bcdefghijklmno":
        sv[f"{w}{i}"].border = Border(left=border_side, top=border_side, bottom=border_side)
    sv[f"o{i}"].border = Border(left=border_side, top=border_side, right=border_side, bottom=border_side)
"""
sv_xl.save(filename="сводка.xlsx")

from aspose.cells import Workbook, PageOrientationType, PaperSizeType
workbook = Workbook("сводка.xlsx")
worksheet = workbook.worksheets.get(0)

worksheet.page_setup.paper_size = PaperSizeType.PAPER_A4
worksheet.page_setup.orientation = PageOrientationType.LANDSCAPE
worksheet.page_setup.top_margin = 0.0
worksheet.page_setup.left_margin = 0.0
worksheet.page_setup.right_margin = 0.0
worksheet.page_setup.bottom_margin = 0.0
worksheet.page_setup.zoom = 90
worksheet.page_setup.center_vertically = True
worksheet.page_setup.center_horizontally = True

#import os
#
#path = fr"C:\Users\SERGO\Documents\апрель"
#os.makedirs(path)
#date = 25.04
#workbook.save(fr"C:\Users\SERGO\Documents\апрель\Сводка{date}.jpg")

workbook.save(fr"Сводка.jpg")